"""Text extraction from PDF / DOCX / TXT resumes (reused and hardened from the original app)."""
import io
import logging
import re

import docx
import pdfplumber

from ..errors import ApiError

log = logging.getLogger(__name__)
MAX_PDF_PAGES = 15


class ExtractionError(ApiError):
    status, code = 422, "unreadable_resume"


def _clean(text):
    text = text.replace("\x00", " ").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(data: bytes, ext: str) -> str:
    try:
        if ext == "pdf":
            pages = []
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                for page in pdf.pages[:MAX_PDF_PAGES]:
                    page_text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
                    if page_text.strip():
                        pages.append(page_text)
            text = "\n\n".join(pages)
        elif ext == "docx":
            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            # Many resumes put content in tables.
            for table in document.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(dict.fromkeys(cells)))
            text = "\n".join(parts)
        elif ext == "txt":
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                text = data.decode("latin-1")
        else:
            raise ExtractionError("Unsupported file type.")
    except ExtractionError:
        raise
    except Exception as exc:
        log.warning("text_extraction_failed", extra={"file_type": ext, "error_type": type(exc).__name__})
        raise ExtractionError("We couldn't read this file. Please upload a text-based PDF, DOCX or TXT.") from exc

    text = _clean(text)
    if len(text) < 50:
        raise ExtractionError(
            "This resume contains too little readable text. Scanned/image PDFs aren't supported - "
            "please upload a text-based PDF, DOCX or TXT.")
    return text
